from __future__ import annotations

from copy import copy
from typing import Any

from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string


def col_to_index(col: str) -> int:
    return column_index_from_string(col)


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def normalize_key(value: Any) -> str | None:
    """
    照合キー用の正規化。
    企業No は文字列比較前提で、前後空白だけ落とす。
    """
    if value is None:
        return None

    text = str(value).strip()
    if text == "":
        return None
    return text


def get_cell_value(ws, row: int, col: str):
    return ws.cell(row=row, column=col_to_index(col)).value


def set_cell_value(ws, row: int, col: str, value) -> None:
    ws.cell(row=row, column=col_to_index(col)).value = value


def _row_has_value(ws, row: int) -> bool:
    for col in range(1, ws.max_column + 1):
        if not is_blank(ws.cell(row=row, column=col).value):
            return True
    return False


def iter_non_blank_rows(ws, start_row: int):
    for row in range(start_row, ws.max_row + 1):
        if _row_has_value(ws, row):
            yield row


def find_last_data_row(ws, start_row: int) -> int:
    for row in range(ws.max_row, start_row - 1, -1):
        if _row_has_value(ws, row):
            return row
    return start_row - 1


def build_row_index_by_key(ws, key_col: str, start_row: int) -> dict[str, int]:
    """
    指定列をキーにして、キー -> 行番号 の辞書を作る。
    同じキーが複数ある場合は、後ろの行を優先する。
    """
    index: dict[str, int] = {}
    key_col_idx = col_to_index(key_col)

    for row in range(start_row, ws.max_row + 1):
        raw_value = ws.cell(row=row, column=key_col_idx).value
        key = normalize_key(raw_value)
        if key is None:
            continue
        index[key] = row

    return index


def next_sequential_number(ws, target_col: str, current_append_row: int, start_number: int = 1) -> int:
    col_idx = col_to_index(target_col)

    for row in range(current_append_row - 1, 0, -1):
        value = ws.cell(row=row, column=col_idx).value

        if isinstance(value, int):
            return value + 1

        if isinstance(value, float):
            return int(value) + 1

        if isinstance(value, str):
            stripped = value.strip()
            if stripped.isdigit():
                return int(stripped) + 1

    return start_number


def copy_row_style_and_formulas(ws, source_row: int, target_row: int) -> None:
    """
    既存行の書式と数式を新規行へコピーする。
    値セルは None にして、あとから processor 側で埋める。
    """
    if source_row <= 0 or target_row <= 0:
        return

    if source_row == target_row:
        return

    source_dim = ws.row_dimensions[source_row]
    target_dim = ws.row_dimensions[target_row]

    if source_dim.height is not None:
        target_dim.height = source_dim.height

    for col_idx in range(1, ws.max_column + 1):
        source_cell = ws.cell(row=source_row, column=col_idx)
        target_cell = ws.cell(row=target_row, column=col_idx)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)

        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format

        if source_cell.font:
            target_cell.font = copy(source_cell.font)

        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)

        if source_cell.border:
            target_cell.border = copy(source_cell.border)

        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)

        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)

        source_value = source_cell.value

        if isinstance(source_value, str) and source_value.startswith("="):
            try:
                target_cell.value = Translator(
                    source_value,
                    origin=source_cell.coordinate,
                ).translate_formula(target_cell.coordinate)
            except Exception:
                target_cell.value = source_value
        else:
            target_cell.value = None