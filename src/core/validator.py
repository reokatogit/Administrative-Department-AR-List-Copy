from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import warnings

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.workbook import Workbook

from core.constants import DATA_START_ROW, HEADER_ROW, MASTER_SHEET_NAME
from core.mapping import MappingConfig


warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
    category=UserWarning,
)


class ValidationError(Exception):
    pass


@dataclass(frozen=True)
class ValidatedWorkbooks:
    monthly_path: Path
    master_path: Path
    output_dir: Path
    monthly_workbook: Workbook
    master_workbook: Workbook


def _validate_xlsx_file(path: str | Path, label: str) -> Path:
    path = Path(path)
    if not path.exists():
        raise ValidationError(f"{label}が見つかりません: {path}")
    if path.suffix.lower() != ".xlsx":
        raise ValidationError(f"{label}は .xlsx ファイルを選択してください。")
    return path


def _validate_output_dir(path: str | Path) -> Path:
    path = Path(path)
    if not path.exists():
        raise ValidationError(f"出力先フォルダが見つかりません: {path}")
    if not path.is_dir():
        raise ValidationError(f"出力先にフォルダを指定してください: {path}")
    return path


def _load_workbook(path: Path, label: str) -> Workbook:
    try:
        return load_workbook(path, data_only=False)
    except Exception as exc:  # noqa: BLE001
        raise ValidationError(f"{label}を読み込めませんでした: {path}\n{exc}") from exc


def _validate_monthly_workbook(workbook: Workbook) -> None:
    if len(workbook.sheetnames) != 1:
        raise ValidationError("月次ファイルはシート1枚のみを想定しています。")

    ws = workbook.worksheets[0]
    if ws.max_row < HEADER_ROW:
        raise ValidationError("月次ファイルのヘッダ行が不足しています。")
    if ws.max_row < DATA_START_ROW:
        raise ValidationError("月次ファイルにデータ行がありません。")


def _validate_master_workbook(workbook: Workbook) -> None:
    if not workbook.sheetnames:
        raise ValidationError("マスターファイルにシートが存在しません。")
    if workbook.sheetnames[0] != MASTER_SHEET_NAME:
        raise ValidationError(f"マスターの先頭シート名は '{MASTER_SHEET_NAME}' である必要があります。")


def _validate_required_columns(workbook: Workbook, required_columns: set[str], workbook_label: str) -> None:
    if not required_columns:
        return

    ws = workbook.worksheets[0]
    max_required_index = max(column_index_from_string(col) for col in required_columns)
    if ws.max_column < max_required_index:
        missing = sorted(required_columns, key=column_index_from_string)
        raise ValidationError(
            f"{workbook_label}の必要列が不足しています。必要列: {', '.join(missing)}"
        )


def validate_inputs(
    monthly_path: str | Path,
    master_path: str | Path,
    output_dir: str | Path,
    mapping: MappingConfig,
) -> ValidatedWorkbooks:
    monthly_path = _validate_xlsx_file(monthly_path, "月次ファイル")
    master_path = _validate_xlsx_file(master_path, "マスターファイル")
    output_dir = _validate_output_dir(output_dir)

    monthly_wb = _load_workbook(monthly_path, "月次ファイル")
    master_wb = _load_workbook(master_path, "マスターファイル")

    _validate_monthly_workbook(monthly_wb)
    _validate_master_workbook(master_wb)
    _validate_required_columns(monthly_wb, mapping.required_monthly_columns(), "月次ファイル")
    _validate_required_columns(master_wb, mapping.required_master_columns(), "マスターファイル")

    return ValidatedWorkbooks(
        monthly_path=monthly_path,
        master_path=master_path,
        output_dir=output_dir,
        monthly_workbook=monthly_wb,
        master_workbook=master_wb,
    )