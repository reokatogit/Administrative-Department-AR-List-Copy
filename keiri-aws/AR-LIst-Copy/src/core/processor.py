from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from openpyxl.utils import get_column_letter

from core.constants import DATA_START_ROW, HEADER_ROW
from core.mapping import ConditionalCopyRule, CopyRule, DEFAULT_MAPPING, MappingConfig
from core.validator import validate_inputs
from utils.excel_helpers import (
    build_row_index_by_key,
    col_to_index,
    copy_row_style_and_formulas,
    find_last_data_row,
    get_cell_value,
    is_blank,
    iter_non_blank_rows,
    next_sequential_number,
    normalize_key,
)


MONTHLY_KEY_COL = "A"
MASTER_KEY_COL = "E"


class ProcessingError(Exception):
    pass


@dataclass(frozen=True)
class ProcessingResult:
    output_path: Path
    source_row_count: int
    updated_row_count: int
    appended_row_count: int


class WorkbookProcessor:
    def __init__(self, mapping: MappingConfig | None = None) -> None:
        self.mapping = mapping or DEFAULT_MAPPING

    def process(
        self,
        monthly_path: str | Path,
        master_path: str | Path,
        output_path: str | Path,
        progress_callback: Callable[[int, str], None] | None = None,
    ) -> ProcessingResult:
        if not self.mapping.is_configured:
            raise ProcessingError("マッピング定義が未設定です。src/core/mapping.py を埋めてください。")

        output_path = Path(output_path)
        validated = validate_inputs(monthly_path, master_path, output_path.parent, self.mapping)

        source_ws = validated.monthly_workbook.worksheets[0]
        target_ws = validated.master_workbook.worksheets[0]

        source_rows = list(iter_non_blank_rows(source_ws, DATA_START_ROW))
        if not source_rows:
            raise ProcessingError("月次ファイルに追記対象のデータ行がありません。")

        total_rows = len(source_rows)
        self._report_progress(progress_callback, 0, "処理を開始しています…")

        key_to_row = build_row_index_by_key(target_ws, MASTER_KEY_COL, DATA_START_ROW)
        append_row = find_last_data_row(target_ws, DATA_START_ROW) + 1

        updated_count = 0
        appended_count = 0
        progress_step = max(1, total_rows // 100)

        for i, source_row in enumerate(source_rows, start=1):
            source_key = normalize_key(get_cell_value(source_ws, source_row, MONTHLY_KEY_COL))

            if source_key and source_key in key_to_row:
                target_row = key_to_row[source_key]
                self._update_existing_row(source_ws, target_ws, source_row, target_row)
                updated_count += 1
            else:
                target_row = append_row
                self._append_new_row(source_ws, target_ws, source_row, target_row)
                appended_count += 1
                append_row += 1

                # 新規追加したら index にも反映
                appended_key = normalize_key(get_cell_value(target_ws, target_row, MASTER_KEY_COL))
                if appended_key:
                    key_to_row[appended_key] = target_row

            if i % progress_step == 0 or i == total_rows:
                percent = int((i / total_rows) * 90)
                self._report_progress(
                    progress_callback,
                    percent,
                    f"処理中… {i}/{total_rows}（更新 {updated_count} 件 / 追加 {appended_count} 件）",
                )

        self._report_progress(progress_callback, 93, "フィルタを更新しています…")
        self._refresh_auto_filter(target_ws)

        self._enable_recalc_on_open(validated.master_workbook)

        self._report_progress(progress_callback, 95, "数式参照を調整しています…")
        self._sanitize_formula_references(validated.master_workbook)

        self._report_progress(progress_callback, 97, "保存しています…")
        self._save_output_workbook(validated.master_workbook, output_path)

        self._report_progress(progress_callback, 100, "完了しました。")

        return ProcessingResult(
            output_path=output_path,
            source_row_count=total_rows,
            updated_row_count=updated_count,
            appended_row_count=appended_count,
        )

    def _update_existing_row(self, source_ws, target_ws, source_row: int, target_row: int) -> None:
        """
        既存行更新:
        - 転記対象列を更新
        - 派生列を更新
        - 固定値を反映
        - 採番、書式コピー、数式コピーはしない
        """
        self._apply_copy_rules(source_ws, target_ws, source_row, target_row)
        self._apply_conditional_copy_rules(source_ws, target_ws, source_row, target_row)
        self._apply_value_map_copy_rules(source_ws, target_ws, source_row, target_row)
        self._apply_fixed_value_rules(target_ws, target_row)

    def _append_new_row(self, source_ws, target_ws, source_row: int, target_row: int) -> None:
        """
        新規追加:
        - 前行から書式/数式コピー
        - 転記
        - 派生列
        - 採番
        - 固定値
        """
        template_row = target_row - 1 if target_row > DATA_START_ROW else DATA_START_ROW
        copy_row_style_and_formulas(target_ws, template_row, target_row)

        self._apply_copy_rules(source_ws, target_ws, source_row, target_row)
        self._apply_conditional_copy_rules(source_ws, target_ws, source_row, target_row)
        self._apply_value_map_copy_rules(source_ws, target_ws, source_row, target_row)
        self._apply_auto_number_rules(target_ws, target_row)
        self._apply_fixed_value_rules(target_ws, target_row)

    def _apply_copy_rules(self, source_ws, target_ws, source_row: int, target_row: int) -> None:
        for rule in self.mapping.always_copy_rules:
            self._copy_value(rule, source_ws, target_ws, source_row, target_row)

    def _apply_conditional_copy_rules(self, source_ws, target_ws, source_row: int, target_row: int) -> None:
        for rule in self.mapping.conditional_copy_rules:
            condition_value = source_ws.cell(
                row=source_row,
                column=col_to_index(rule.condition_source_col),
            ).value
            if condition_value in rule.allowed_values:
                self._copy_value(rule, source_ws, target_ws, source_row, target_row)

    def _apply_value_map_copy_rules(self, source_ws, target_ws, source_row: int, target_row: int) -> None:
        for rule in self.mapping.value_map_copy_rules:
            source_value = source_ws.cell(
                row=source_row,
                column=col_to_index(rule.source_col),
            ).value
            mapped_value = rule.value_map.get(source_value, rule.default_value)
            target_ws.cell(
                row=target_row,
                column=col_to_index(rule.target_col),
            ).value = mapped_value

    def _apply_fixed_value_rules(self, target_ws, target_row: int) -> None:
        for rule in self.mapping.fixed_value_rules:
            target_cell = target_ws.cell(
                row=target_row,
                column=col_to_index(rule.target_col),
            )
            if rule.only_if_blank and not is_blank(target_cell.value):
                continue
            target_cell.value = rule.value

    def _apply_auto_number_rules(self, target_ws, target_row: int) -> None:
        for rule in self.mapping.auto_number_rules:
            target_ws.cell(
                row=target_row,
                column=col_to_index(rule.target_col),
            ).value = next_sequential_number(
                ws=target_ws,
                target_col=rule.target_col,
                current_append_row=target_row,
                start_number=rule.start_number,
            )

    @staticmethod
    def _copy_value(
        rule: CopyRule | ConditionalCopyRule,
        source_ws,
        target_ws,
        source_row: int,
        target_row: int,
    ) -> None:
        source_value = source_ws.cell(
            row=source_row,
            column=col_to_index(rule.source_col),
        ).value
        target_ws.cell(
            row=target_row,
            column=col_to_index(rule.target_col),
        ).value = source_value

    @staticmethod
    def _refresh_auto_filter(target_ws) -> None:
        if not target_ws.auto_filter:
            return

        last_row = find_last_data_row(target_ws, DATA_START_ROW)
        last_col = get_column_letter(target_ws.max_column)
        target_ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{last_row}"

    @staticmethod
    def _enable_recalc_on_open(workbook) -> None:
        calculation = getattr(workbook, "calculation", None)
        if calculation is None:
            return
        if hasattr(calculation, "fullCalcOnLoad"):
            calculation.fullCalcOnLoad = True
        if hasattr(calculation, "forceFullCalc"):
            calculation.forceFullCalc = True

    @staticmethod
    def _save_output_workbook(workbook, output_path: Path) -> None:
        try:
            workbook.save(output_path)
        except Exception as exc:  # noqa: BLE001
            raise ProcessingError(f"出力ファイルの保存に失敗しました。\n{output_path}\n{exc}") from exc

    @staticmethod
    def _report_progress(
        progress_callback: Callable[[int, str], None] | None,
        percent: int,
        message: str,
    ) -> None:
        if progress_callback is None:
            return

        try:
            progress_callback(percent, message)
        except Exception:
            pass

    @staticmethod
    def _sanitize_formula_references(workbook) -> None:
        """
        [1]ALL!E:E のような疑似外部参照を、内部参照 ALL!E:E に戻す。
        重くなりすぎないよう、実際に値を持っているセルだけを見る。
        """
        bracketed_ref_pattern = re.compile(r"\[(\d+)\](?=[^!]+!)")

        for ws in workbook.worksheets:
            cells = getattr(ws, "_cells", None)
            if not cells:
                continue

            for cell in cells.values():
                value = cell.value

                if not isinstance(value, str):
                    continue
                if not value.startswith("="):
                    continue
                if "[" not in value or "]" not in value:
                    continue

                normalized = bracketed_ref_pattern.sub("", value)

                if normalized != value:
                    cell.value = normalized