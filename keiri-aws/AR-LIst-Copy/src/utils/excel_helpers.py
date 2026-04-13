from __future__ import annotations

from copy import copy

from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string


def col_to_index(col: str) -> int:
    return column_index_from_string(col)


def is_blank(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def _row_has_value(ws, row: int) -> bool:
    for col in range(1, ws.max_column + 1):
        if not is_blank(ws.cell(row=row, column=col).value):
            return True
    return False


def iter_non_blank_rows(ws, start_row: int):
    for row in range(start_row, ws.max_row + 1):
        if _row_has_value(ws, row):
            yield row


def find_last_data_row(ws, start_row: int) -> int:
    for row in range(ws.max_row, start_row - 1, -1):
        if _row_has_value(ws, row):
            return row
    return start_row - 1


def next_sequential_number(ws, target_col: str, current_append_row: int, start_number: int = 1) -> int:
    col_idx = col_to_index(target_col)

    for row in range(current_append_row - 1, 0, -1):
        value = ws.cell(row=row, column=col_idx).value

        if isinstance(value, int):
            return value + 1

        if isinstance(value, float):
            return int(value) + 1

        if isinstance(value, str):
            stripped = value.strip()
            if stripped.isdigit():
                return int(stripped) + 1

    return start_number


def copy_row_style_and_formulas(ws, source_row: int, target_row: int) -> None:
    if source_row <= 0 or target_row <= 0:
        return

    if source_row == target_row:
        return

    source_dim = ws.row_dimensions[source_row]
    target_dim = ws.row_dimensions[target_row]

    if source_dim.height is not None:
        target_dim.height = source_dim.height

    for col_idx in range(1, ws.max_column + 1):
        source_cell = ws.cell(row=source_row, column=col_idx)
        target_cell = ws.cell(row=target_row, column=col_idx)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)

        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format

        if source_cell.font:
            target_cell.font = copy(source_cell.font)

        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)

        if source_cell.border:
            target_cell.border = copy(source_cell.border)

        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)

        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)

        source_value = source_cell.value

        if isinstance(source_value, str) and source_value.startswith("="):
            try:
                target_cell.value = Translator(
                    source_value,
                    origin=source_cell.coordinate,
                ).translate_formula(target_cell.coordinate)
            except Exception:
                target_cell.value = source_value
        else:
            target_cell.value = None